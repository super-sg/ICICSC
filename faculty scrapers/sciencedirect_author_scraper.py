"""
ScienceDirect Author Metadata Automation Engine
================================================
A resilient, fault-tolerant Playwright automation script designed to extract
academic author metadata across conference technical tracks from ScienceDirect.

Features:
- Cloudflare & Bot-Mitigation: Persistent browser context with stealth launch flags
- Explicit Waits: Uses wait_for_selector(..., state='visible')
- Interactive Flyout Handling: Sequentially triggers author slide drawers (#side-panel)
- Domestic Author Filtering: Filters out Indian authors, retaining only foreign researchers
- Quality Filtering: Retains only records with confirmed email or valid institutional affiliation
- Multi-Tab Professional Excel Export: Navy header, zebra rows, freeze panes, auto-fit columns
- Modularity: Supports specific topics, entire tracks, or the full 7-track conference catalog

Author: Senior Automation Engineer
"""

import os
import re
import sys
import time
import random
import asyncio
import logging
import argparse
from typing import List, Dict, Any, Optional
from urllib.parse import quote_plus

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from playwright.async_api import async_playwright, BrowserContext, Page, TimeoutError as PlaywrightTimeoutError

# ==============================================================================
# LOGGING CONFIGURATION
# ==============================================================================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger("ScienceDirectScraper")

# ==============================================================================
# CONFERENCE TRACKS & TOPICS REGISTRY (ICICSC 2027)
# ==============================================================================
CONFERENCE_TRACKS: Dict[str, Dict[str, Any]] = {
    "Track 1": {
        "name": "Artificial Intelligence & Machine Learning",
        "topics": [
            "Deep Learning Architectures and Applications",
            "Generative AI and Large Language Models (LLMs)",
            "Reinforcement Learning and Decision Systems",
            "Natural Language Processing and Understanding",
            "Explainable, Trustworthy and Responsible AI",
            "Federated and Distributed Learning",
            "AI for Healthcare Diagnostics and Biomedical Systems",
            "Neural Network Optimization Techniques",
            "Transfer Learning and Few-Shot Learning",
            "Multi-Agent and Agentic AI Systems"
        ]
    },
    "Track 2": {
        "name": "Internet of Things & Cyber-Physical Systems",
        "topics": [
            "IoT Architectures, Protocols and Standards",
            "Industrial IoT (IIoT) for Smart Manufacturing",
            "Smart Cities and Urban Infrastructure",
            "Wearable and Ubiquitous Computing",
            "IoT Security, Privacy and Device Authentication",
            "Edge and Fog Computing for IoT",
            "Digital Twins for Cyber-Physical Systems",
            "Precision Agriculture and Agri-Tech using IoT",
            "Smart Healthcare and Remote Patient Monitoring",
            "Energy-Efficient and Low-Power IoT Networks"
        ]
    },
    "Track 3": {
        "name": "Robotics & Intelligent Automation",
        "topics": [
            "Autonomous Mobile Robots and Navigation",
            "Human-Robot Interaction and Collaboration",
            "Swarm Robotics and Multi-Robot Systems",
            "Robotic Process Automation (RPA)",
            "Robot Perception, Sensing and SLAM",
            "Industrial Robotics and Industry 5.0",
            "Medical and Surgical Robotics",
            "Robotic Manipulation, Grasping and Haptics",
            "Bio-Inspired and Soft Robotics",
            "Unmanned Aerial Vehicles (UAV) and Autonomous Drones"
        ]
    },
    "Track 4": {
        "name": "Data Analytics & Big Data",
        "topics": [
            "Big Data Architectures and Scalable Processing",
            "Predictive and Prescriptive Analytics",
            "Real-Time Stream Data Processing",
            "Data Mining and Pattern Discovery",
            "Business Intelligence and Decision Support Systems",
            "Graph-Based Data Analytics and Network Science",
            "Time-Series Analysis and Forecasting",
            "Data Warehousing, Lakehouses and ETL Pipelines",
            "Social Media and Sentiment Analytics",
            "Analytics for Sustainability and ESG Reporting"
        ]
    },
    "Track 5": {
        "name": "Blockchain & Distributed Ledger Technologies",
        "topics": [
            "Blockchain Architectures and Consensus Mechanisms",
            "Smart Contracts and Decentralized Applications (DApps)",
            "Blockchain for Supply Chain Traceability",
            "Cryptocurrency, Tokenomics and Digital Asset Systems",
            "Decentralized Finance (DeFi) Protocols",
            "Blockchain Interoperability and Cross-Chain Bridges",
            "Blockchain Scalability and Layer-2 Solutions",
            "Blockchain for Secure Healthcare Record Management",
            "Non-Fungible Tokens (NFTs) and Digital Ownership",
            "Energy-Efficient and Green Blockchain Systems"
        ]
    },
    "Track 6": {
        "name": "Cybersecurity & Privacy",
        "topics": [
            "Network Security, Intrusion Detection and Prevention",
            "Malware Analysis and Cyber Threat Intelligence",
            "Cryptography and Post-Quantum Security",
            "Zero Trust Architecture and Identity Management",
            "Digital Forensics and Incident Response",
            "Privacy-Preserving Computation and Differential Privacy",
            "Cloud, Virtualization and Container Security",
            "Cybersecurity for Critical Infrastructure and SCADA",
            "Ethical Hacking and Vulnerability Assessment",
            "AI-Driven Cyber Threat and Ransomware Mitigation"
        ]
    },
    "Track 7": {
        "name": "Computer Vision, Cloud & Distributed Systems",
        "topics": [
            "Object Detection, Recognition and Tracking",
            "Medical Image Analysis, Diagnostics and Segmentation",
            "Facial Recognition and Biometric Vision Systems",
            "Video Analytics and Intelligent Surveillance",
            "Multi-Modal Vision-Language Models",
            "Cloud Service Models, Microservices and Container Orchestration",
            "Serverless and Function-as-a-Service (FaaS) Computing",
            "Cloud-Native Application Development and DevOps",
            "Edge-Cloud Continuum and Distributed Consensus",
            "Cloud Computing for AI/ML Workloads and Green Computing"
        ]
    }
}

# ==============================================================================
# DOMESTIC / INDIAN INSTITUTION FILTER KEYWORDS
# ==============================================================================
DOMESTIC_KEYWORDS = [
    'india', 'indian', 'iit', 'nit', 'iiit', 'bits pilani', 'iisc', 'iim', 'isro', 'drdo',
    'delhi', 'mumbai', 'chennai', 'kolkata', 'bengaluru', 'bangalore', 'hyderabad', 'pune',
    'ahmedabad', 'jaipur', 'lucknow', 'patna', 'kanpur', 'nagpur', 'indore', 'bhopal',
    'visakhapatnam', 'coimbatore', 'kochi', 'thiruvananthapuram', 'roorkee', 'kharagpur',
    'guwahati', 'varanasi', 'surat', 'vellore', 'manipal', 'jadavpur', 'amrita', 'sharda',
    'srm', 'thapar', 'vit', 'anna university', 'osmania', 'calcutta university'
]

# ==============================================================================
# EXCEL STYLING CONSTANTS
# ==============================================================================
EXCEL_COLUMNS = [
    'Track',
    'Topic',
    'Paper Title',
    'Author Name',
    'Email',
    'Institution',
    'Department',
    'Designation',
    'Country',
    'Profile / Article URL'
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")


def is_domestic_affiliation(affiliation_text: str) -> bool:
    """Returns True if affiliation indicates an Indian/domestic institution."""
    if not affiliation_text:
        return False
    text_lower = affiliation_text.lower()
    for kw in DOMESTIC_KEYWORDS:
        if re.search(r'\b' + re.escape(kw) + r'\b', text_lower):
            return True
    return False


def parse_affiliation_details(raw_text: str) -> Dict[str, str]:
    """
    Parses unstructured affiliation string into Department, Institution, Designation, Country.
    """
    if not raw_text or not raw_text.strip():
        return {'Department': 'N/A', 'Institution': 'N/A', 'Designation': 'N/A', 'Country': 'N/A'}
    
    parts = [p.strip() for p in raw_text.split(',') if p.strip()]
    if not parts:
        return {'Department': 'N/A', 'Institution': raw_text.strip(), 'Designation': 'N/A', 'Country': 'N/A'}
    
    # Country is usually the last comma segment
    country = parts[-1]
    # Strip state abbreviations and postal codes
    country = re.sub(r'^[A-Z]{2}\s+\d+[\s\w-]*', '', country).strip()
    if re.match(r'^\d+$', country) and len(parts) > 1:
        country = parts[-2]
        
    department = 'N/A'
    institution = 'N/A'
    designation = 'N/A'
    
    desig_patterns = [
        r'\bprofessor\b', r'\bassociate professor\b', r'\bassistant professor\b',
        r'\bresearcher\b', r'\bpostdoc\b', r'\bscientist\b', r'\bfellow\b', r'\bphd candidate\b'
    ]
    for p in parts:
        if any(re.search(pat, p, re.IGNORECASE) for pat in desig_patterns):
            designation = p
            break
            
    dept_patterns = [
        r'\bdepartment\b', r'\bschool\b', r'\bfaculty\b', r'\bdivision\b',
        r'\bcenter\b', r'\bcentre\b', r'\blaboratory\b', r'\blab\b', r'\bdiscipline\b'
    ]
    for p in parts:
        if any(re.search(pat, p, re.IGNORECASE) for pat in dept_patterns):
            department = p
            break
            
    inst_patterns = [
        r'\buniversity\b', r'\binstitute\b', r'\bcollege\b', r'\bacade\w+\b',
        r'\bpolytechnic\b', r'\bhospital\b', r'\bcorporation\b', r'\binc\b', r'\bltd\b'
    ]
    for p in parts:
        if p != department and any(re.search(pat, p, re.IGNORECASE) for pat in inst_patterns):
            institution = p
            break
            
    if institution == 'N/A':
        for p in parts:
            if p != department and p != country and p != designation:
                institution = p
                break

    return {
        'Department': department,
        'Institution': institution,
        'Designation': designation,
        'Country': country
    }


class ScienceDirectAuthorScraper:
    """
    Robust Playwright automation engine for ScienceDirect author metadata extraction.
    """

    def __init__(
        self,
        user_data_dir: Optional[str] = None,
        headless: bool = False,
        delay_min: float = 1.0,
        delay_max: float = 2.5,
        timeout_ms: int = 30000
    ):
        self.user_data_dir = user_data_dir or os.path.expanduser('~/.sd_browser_profile')
        self.headless = headless
        self.delay_min = delay_min
        self.delay_max = delay_max
        self.timeout_ms = timeout_ms
        self.playwright = None
        self.context: Optional[BrowserContext] = None
        self.page: Optional[Page] = None
        self.extracted_records: List[Dict[str, Any]] = []

    async def __aenter__(self):
        await self.initialize()
        return self

    async def __aexit__(self, exc_type, exc_val, exc_tb):
        await self.shutdown()

    async def initialize(self):
        """Launches persistent context with anti-bot evasion arguments."""
        logger.info(f"Initializing browser engine (User Data: {self.user_data_dir}, Headless: {self.headless})...")
        os.makedirs(self.user_data_dir, exist_ok=True)
        self.playwright = await async_playwright().start()
        
        self.context = await self.playwright.chromium.launch_persistent_context(
            user_data_dir=self.user_data_dir,
            channel='chrome',
            headless=self.headless,
            viewport={'width': 1440, 'height': 900},
            locale='en-US',
            timezone_id='Asia/Kolkata',
            args=[
                '--disable-blink-features=AutomationControlled',
                '--no-sandbox',
                '--disable-setuid-sandbox',
                '--disable-infobars',
                '--window-position=0,0'
            ]
        )
        
        # Override navigator.webdriver
        await self.context.add_init_script("""
            Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
            window.chrome = { runtime: {} };
        """)
        
        self.page = self.context.pages[0] if self.context.pages else await self.context.new_page()
        self.page.set_default_timeout(self.timeout_ms)
        logger.info("Browser initialized successfully.")

    async def shutdown(self):
        """Gracefully closes browser and Playwright process."""
        logger.info("Closing browser session...")
        try:
            if self.context:
                await self.context.close()
            if self.playwright:
                await self.playwright.stop()
        except Exception as e:
            logger.warning(f"Error during shutdown: {e}")

    async def polite_delay(self):
        """Randomized humanized delay between actions."""
        pause = random.uniform(self.delay_min, self.delay_max)
        await self.page.wait_for_timeout(pause * 1000)

    async def search_topic(self, topic: str, max_pages: int = 2) -> List[str]:
        """
        Navigates search results for a given query and collects article links.
        ScienceDirect uses offset=0, offset=25, offset=50 for pagination.
        """
        logger.info(f"Initiating ScienceDirect search for query: '{topic}' (Pages: {max_pages})")
        article_links: List[str] = []
        seen_urls = set()

        for page_idx in range(max_pages):
            offset = page_idx * 25
            encoded_query = quote_plus(topic)
            url = f"https://www.sciencedirect.com/search?qs={encoded_query}"
            if offset > 0:
                url += f"&offset={offset}"

            logger.info(f"Navigating to search page {page_idx + 1}/{max_pages}: {url}")
            try:
                response = await self.page.goto(url, wait_until='networkidle')
                if response and response.status == 403:
                    logger.warning("Encountered 403 Challenge. Waiting 5s for persistent profile resolution...")
                    await self.page.wait_for_timeout(5000)

                # Wait for search results or article links to appear
                await self.page.wait_for_selector('a[href*="/science/article/pii/"]', timeout=15000)
                await self.polite_delay()

                links = await self.page.eval_on_selector_all(
                    'a[href*="/science/article/pii/"]',
                    'els => els.map(e => e.href)'
                )

                for link in links:
                    clean_link = link.split('?')[0].split('#')[0]
                    if '/pdfft' in clean_link:
                        continue
                    if clean_link not in seen_urls and '/science/article/pii/' in clean_link:
                        seen_urls.add(clean_link)
                        article_links.append(clean_link)

                logger.info(f"Page {page_idx + 1}: Found {len(links)} article elements ({len(article_links)} unique accumulated)")

            except PlaywrightTimeoutError:
                logger.warning(f"Search results timed out on page {page_idx + 1} for '{topic}'. Proceeding with collected links.")
                break
            except Exception as e:
                logger.error(f"Error scraping search page {page_idx + 1} for '{topic}': {e}")
                break

        return article_links

    async def scrape_article_authors(
        self,
        article_url: str,
        track_name: str,
        topic_name: str
    ) -> List[Dict[str, Any]]:
        """
        Navigates to an article, triggers side drawer for each author, and extracts metadata.
        """
        logger.info(f"Scraping article: {article_url}")
        records: List[Dict[str, Any]] = []

        try:
            await self.page.goto(article_url, wait_until='networkidle')
            await self.polite_delay()
        except Exception as e:
            logger.warning(f"Initial navigation slow for {article_url}: {e}")
            try:
                await self.page.wait_for_selector('.content-title, h1', timeout=10000)
            except Exception:
                logger.error(f"Failed to load article {article_url}. Skipping.")
                return records

        # 1. Extract Paper Title
        title = "N/A"
        try:
            title_el = await self.page.query_selector('.content-title, h1.article-title, #title-content, h1')
            if title_el:
                title = await title_el.inner_text()
                title = title.strip().replace('\n', ' ')
        except Exception as e:
            logger.debug(f"Title extraction error: {e}")

        # 2. Locate Author Trigger Buttons
        author_btn_selector = (
            '#author-group button[data-sd-ui-side-panel-opener="true"], '
            'button[data-sd-ui-side-panel-opener="true"][data-xocs-content-type="author"], '
            'button:has(.given-name), button:has(.surname), button:has(.react-xocs-alternative-link)'
        )
        try:
            author_btns = await self.page.query_selector_all(author_btn_selector)
        except Exception:
            author_btns = []

        if not author_btns:
            logger.info(f"No interactive author buttons found on {article_url}")
            return records

        logger.info(f"Found {len(author_btns)} author trigger(s) on paper.")

        # 3. Interactively Click Each Author Drawer
        for idx in range(len(author_btns)):
            try:
                # Re-query elements to guard against DOM re-renders
                btns = await self.page.query_selector_all(author_btn_selector)
                if idx >= len(btns):
                    break
                target_btn = btns[idx]

                # Scroll into view and click
                await target_btn.scroll_into_view_if_needed()
                await self.polite_delay()
                await target_btn.click()

                # Explicit wait for side panel
                try:
                    await self.page.wait_for_selector('#side-panel', state='visible', timeout=5000)
                except PlaywrightTimeoutError:
                    # Attempt click fallback
                    await target_btn.dispatch_event('click')
                    await self.page.wait_for_selector('#side-panel', state='visible', timeout=5000)

                await self.page.wait_for_timeout(800)

                # Extract details from #side-panel
                panel = await self.page.query_selector('#side-panel')
                if not panel:
                    continue

                # Name
                name_el = await panel.query_selector('.side-panel-author h2, h2')
                author_name = await name_el.inner_text() if name_el else "N/A"
                author_name = author_name.strip()

                # Email
                mailto_el = await panel.query_selector('a[href^="mailto:"]')
                email = "N/A"
                if mailto_el:
                    href = await mailto_el.get_attribute('href')
                    if href:
                        email = href.replace('mailto:', '').split('?')[0].strip()

                # Affiliations
                affil_els = await panel.query_selector_all('.affiliation')
                raw_affils = []
                for af in affil_els:
                    txt = await af.inner_text()
                    txt = txt.strip()
                    if txt and txt not in raw_affils:
                        raw_affils.append(txt)
                full_affiliation = " | ".join(raw_affils) if raw_affils else "N/A"

                # Parse affiliation fields
                parsed = parse_affiliation_details(raw_affils[0] if raw_affils else "")

                # Close the side panel to reset state for next author
                close_btn = await panel.query_selector('.side-panel-close-btn, button[aria-label="Close Author panel"]')
                if close_btn:
                    await close_btn.click()
                    await self.page.wait_for_timeout(500)
                else:
                    # Press Escape as fallback
                    await self.page.keyboard.press("Escape")
                    await self.page.wait_for_timeout(500)

                # 4. Apply Filters:
                # - Filter out domestic/Indian authors
                # - Must have valid email OR institution
                is_indian = is_domestic_affiliation(full_affiliation)
                has_valid_info = (email != "N/A") or (parsed['Institution'] != "N/A")

                if is_indian:
                    logger.debug(f"Filtered DOMESTIC author: {author_name} ({full_affiliation[:40]}...)")
                    continue

                if not has_valid_info:
                    logger.debug(f"Filtered author with no email and no institution: {author_name}")
                    continue

                record = {
                    'Track': track_name,
                    'Topic': topic_name,
                    'Paper Title': title,
                    'Author Name': author_name,
                    'Email': email,
                    'Institution': parsed['Institution'],
                    'Department': parsed['Department'],
                    'Designation': parsed['Designation'],
                    'Country': parsed['Country'],
                    'Profile / Article URL': article_url
                }
                records.append(record)
                logger.info(f"  ✓ Extracted Foreign Author: {author_name} | {parsed['Country']} | {email} | {parsed['Institution'][:30]}")

            except Exception as e:
                logger.warning(f"Error extracting author {idx} on {article_url}: {e}")
                # Ensure panel is dismissed before next author
                try:
                    await self.page.keyboard.press("Escape")
                except Exception:
                    pass

        return records

    async def execute_topics(
        self,
        topic_queue: List[Dict[str, str]],
        max_pages: int = 2,
        max_articles_per_topic: int = 15
    ) -> List[Dict[str, Any]]:
        """
        Iterates through a queue of topics, scrapes articles, and collects foreign author records.
        """
        all_results: List[Dict[str, Any]] = []

        for item in topic_queue:
            track_name = item['track']
            topic_name = item['topic']
            logger.info(f"\n=======================================================")
            logger.info(f"PROCESSING TOPIC: [{track_name}] {topic_name}")
            logger.info(f"=======================================================")

            links = await self.search_topic(topic_name, max_pages=max_pages)
            if not links:
                logger.warning(f"No articles retrieved for '{topic_name}'.")
                continue

            target_links = links[:max_articles_per_topic]
            logger.info(f"Proceeding to scrape top {len(target_links)} articles for '{topic_name}'...")

            for a_idx, link in enumerate(target_links, 1):
                logger.info(f"[{a_idx}/{len(target_links)}] Opening paper: {link}")
                authors = await self.scrape_article_authors(link, track_name, topic_name)
                all_results.extend(authors)
                await self.polite_delay()

        self.extracted_records = all_results
        return all_results


def export_to_excel(
    data: List[Dict[str, Any]],
    output_filepath: str,
    group_by_track: bool = True
) -> None:
    """
    Exports extracted author records to a styled multi-tab Excel file.
    """
    if not data:
        logger.warning("No records to export.")
        return

    df = pd.DataFrame(data)
    for col in EXCEL_COLUMNS:
        if col not in df.columns:
            df[col] = "N/A"
    df = df[EXCEL_COLUMNS].fillna("N/A")

    logger.info(f"Exporting {len(df)} records to {output_filepath}...")
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    def style_ws(ws):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 26

        for col_idx, cell in enumerate(ws[1], 1):
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')

        col_widths = {c: len(str(ws.cell(row=1, column=c).value or '')) for c in range(1, len(EXCEL_COLUMNS) + 1)}

        for r_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
            ws.row_dimensions[r_idx].height = 20
            is_even = (r_idx % 2 == 0)
            for c_idx, cell in enumerate(row, 1):
                cell.font = REGULAR_FONT
                cell.border = BORDER_THIN
                if is_even:
                    cell.fill = ZEBRA_FILL

                val = str(cell.value or '')
                col_widths[c_idx] = max(col_widths[c_idx], len(val))
                cell.alignment = Alignment(horizontal='left', vertical='center')

        for c_idx, max_l in col_widths.items():
            col_let = get_column_letter(c_idx)
            ws.column_dimensions[col_let].width = min(max(max_l + 4, 12), 48)

        ws.auto_filter.ref = ws.dimensions

    def write_sheet(ws, sub_df):
        ws.append(EXCEL_COLUMNS)
        for _, row in sub_df.iterrows():
            clean_row = []
            for val in row.tolist():
                if pd.isna(val) or val is None or str(val).strip() in ['', 'nan', 'NaN']:
                    clean_row.append('N/A')
                else:
                    clean_row.append(val)
            ws.append(clean_row)
        style_ws(ws)

    # Master sheet with all records
    master_ws = wb.create_sheet(title="All Foreign Authors")
    write_sheet(master_ws, df)

    # Individual tabs by Track
    if group_by_track and 'Track' in df.columns:
        tracks = df['Track'].dropna().unique()
        for t in tracks:
            t_df = df[df['Track'] == t]
            tab_name = str(t)[:31].replace(':', '').strip()
            t_ws = wb.create_sheet(title=tab_name)
            write_sheet(t_ws, t_df)

    wb.remove(default_sheet)
    wb.save(output_filepath)
    logger.info(f"Successfully generated styled Excel file: {output_filepath} ({len(wb.sheetnames)} tabs)")


# ==============================================================================
# CLI DISPATCHER
# ==============================================================================
def parse_arguments():
    parser = argparse.ArgumentParser(
        description="ScienceDirect Author Extraction Automation Engine"
    )
    parser.add_argument(
        "--track",
        type=str,
        default="1",
        help="Track key ('1', '2', ..., '7', or 'all'). Default: '1'"
    )
    parser.add_argument(
        "--topic",
        type=str,
        default=None,
        help="Specific topic query to search (overrides track list)."
    )
    parser.add_argument(
        "--pages",
        type=int,
        default=1,
        help="Number of search result pages to paginate per topic (default: 1)."
    )
    parser.add_argument(
        "--max-articles",
        type=int,
        default=5,
        help="Maximum articles to scrape per topic (default: 5)."
    )
    parser.add_argument(
        "--output",
        type=str,
        default="ScienceDirect_Foreign_Authors.xlsx",
        help="Output Excel file path."
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run browser in headless mode (headless=False recommended for Cloudflare)."
    )
    return parser.parse_args()


async def main():
    args = parse_arguments()

    # Build work queue
    work_queue: List[Dict[str, str]] = []
    if args.topic:
        work_queue.append({"track": "Custom", "topic": args.topic})
    elif args.track.lower() == "all":
        for t_key, t_data in CONFERENCE_TRACKS.items():
            t_title = f"{t_key}: {t_data['name']}"
            for sub_topic in t_data['topics']:
                work_queue.append({"track": t_title, "topic": sub_topic})
    else:
        t_key = f"Track {args.track}" if not args.track.lower().startswith("track") else args.track
        if t_key in CONFERENCE_TRACKS:
            t_data = CONFERENCE_TRACKS[t_key]
            t_title = f"{t_key}: {t_data['name']}"
            for sub_topic in t_data['topics']:
                work_queue.append({"track": t_title, "topic": sub_topic})
        else:
            print(f"Error: Unknown track '{args.track}'. Choose from 1 to 7 or 'all'.")
            sys.exit(1)

    logger.info(f"Initialized automation run with {len(work_queue)} topic(s) to process.")

    async with ScienceDirectAuthorScraper(headless=args.headless) as scraper:
        records = await scraper.execute_topics(
            topic_queue=work_queue,
            max_pages=args.pages,
            max_articles_per_topic=args.max_articles
        )

        export_to_excel(records, args.output)
        logger.info(f"Scraping run finished. Total foreign authors extracted: {len(records)}")


if __name__ == "__main__":
    asyncio.run(main())
