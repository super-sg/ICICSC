import os
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

STANDARD_COLUMNS = ['Name', 'Institution', 'Department', 'Designation', 'Qualification', 'Email', 'Profile URL']

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

def get_clean_tab_name(fname, df):
    if len(df) > 0 and 'Institution' in df.columns:
        inst = str(df['Institution'].iloc[0]).strip()
        # Clean long names for Excel tabs (max 31 chars, no invalid chars)
        if 'Maulana Azad' in inst: return 'MANIT Bhopal'
        if 'Malaviya' in inst: return 'MNIT Jaipur'
        if 'Motilal Nehru' in inst: return 'MNNIT Allahabad'
        if 'Visvesvaraya' in inst: return 'VNIT Nagpur'
        if 'Sardar Vallabhbhai' in inst or 'SVNIT' in inst: return 'SVNIT Surat'
        if 'IIEST' in inst: return 'IIEST Shibpur'
        if 'National Institute of Technology' in inst:
            clean = inst.replace('National Institute of Technology', 'NIT').strip()
            if len(clean) <= 31: return clean
        if 'Indian Institute of Information Technology' in inst:
            clean = inst.replace('Indian Institute of Information Technology', 'IIIT').strip()
            if len(clean) <= 31: return clean
        if 'K. K. Birla' in inst or 'Goa Campus' in inst: return 'BITS Goa'
        if 'Hyderabad Campus' in inst: return 'BITS Hyderabad'
        if 'Pilani Campus' in inst: return 'BITS Pilani'
        if 'Dubai Campus' in inst: return 'BITS Dubai'
        if 'Main Campus' in inst: return 'BIT Mesra'
        if 'PDPM IIITDM' in inst: return 'IIITDM Jabalpur'
        if 'ABV-IIITM' in inst: return 'IIITM Gwalior'
        if 'NIT Karnataka' in inst: return 'NIT Surathkal'
        if 'NIT Tiruchirappalli' in inst: return 'NIT Trichy'
        if 'IIIT Tiruchirappalli' in inst: return 'IIIT Trichy'
        if len(inst) <= 31 and not any(c in inst for c in ['[', ']', ':', '*', '?', '/', '\\']):
            return inst

    # Fallback from filename
    name = fname.replace('_Faculty.xlsx', '').replace('_Complete', '').replace('_', ' ')
    name = re.sub(r'[\[\]\:\*\?\/\\ ]+', ' ', name).strip()
    return name[:31].strip()

def style_sheet(ws):
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 26
    
    # Header styling
    for col_idx, cell in enumerate(ws[1], 1):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
        
    # Data row styling & column width calculation
    col_widths = {col_idx: len(str(cell.value or '')) for col_idx, cell in enumerate(ws[1], 1)}
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        for col_idx, cell in enumerate(row, 1):
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN
            if is_even:
                cell.fill = ZEBRA_FILL
            
            val_str = str(cell.value or '')
            col_widths[col_idx] = max(col_widths[col_idx], len(val_str))
            
            # Alignments
            col_name = ws.cell(row=1, column=col_idx).value
            if col_name in ['Qualification']:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Set column widths
    for col_idx, max_len in col_widths.items():
        col_letter = get_column_letter(col_idx)
        # Cap reasonable width with padding
        adjusted_width = min(max(max_len + 4, 12), 48)
        ws.column_dimensions[col_letter].width = adjusted_width
        
    # Set Auto-filter
    ws.auto_filter.ref = ws.dimensions

def build_consolidated_workbook(folder_name, output_master_name, output_alias_name):
    folder_path = os.path.join(BASE_DIR, folder_name)
    if not os.path.exists(folder_path):
        print(f"Folder not found: {folder_path}")
        return
        
    # Collect all individual xlsx files (exclude any previously generated master/all files)
    files = sorted([
        f for f in os.listdir(folder_path)
        if f.endswith('.xlsx') and not f.startswith('~') and not f.startswith('.') 
        and not 'Master' in f and not 'All_Campuses' in f and not 'All_Institutions' in f
    ])
    
    print(f"\nProcessing {folder_name}: Found {len(files)} branch files.")
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    
    all_rows = []
    used_tab_names = set()
    
    # Process each branch file into its own tab
    for f in files:
        fpath = os.path.join(folder_path, f)
        df = pd.read_excel(fpath)
        
        # Ensure standard columns
        for col in STANDARD_COLUMNS:
            if col not in df.columns:
                df[col] = "N/A"
        df = df[STANDARD_COLUMNS]
        
        all_rows.append(df)
        
        tab_name = get_clean_tab_name(f, df)
        # Avoid any collisions
        if tab_name in used_tab_names:
            tab_name = f"{tab_name[:28]}_{len(used_tab_names)}"
        used_tab_names.add(tab_name)
        
        ws = wb.create_sheet(title=tab_name)
        ws.append(STANDARD_COLUMNS)
        for _, row in df.iterrows():
            ws.append(row.tolist())
            
        style_sheet(ws)
        print(f"  Added tab: [{tab_name}] with {len(df)} records")

    # Create the 'All Faculty' tab as the FIRST sheet
    all_df = pd.concat(all_rows, ignore_index=True) if all_rows else pd.DataFrame(columns=STANDARD_COLUMNS)
    all_ws = wb.create_sheet(title="All Faculty", index=0)
    all_ws.append(STANDARD_COLUMNS)
    for _, row in all_df.iterrows():
        all_ws.append(row.tolist())
        
    style_sheet(all_ws)
    print(f"  Added Master tab: [All Faculty] with {len(all_df)} consolidated records")
    
    # Remove initial blank sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)
        
    # Save master workbook inside the folder
    out_master_path = os.path.join(folder_path, output_master_name)
    wb.save(out_master_path)
    print(f"  [Saved Master] {out_master_path}")
    
    # Also save with alias name in the folder
    if output_alias_name:
        out_alias_path = os.path.join(folder_path, output_alias_name)
        wb.save(out_alias_path)
        print(f"  [Saved Alias]  {out_alias_path}")

def main():
    print("=== Generating Consolidated Multi-Tab Workbooks ===")
    
    # 1. BIT
    build_consolidated_workbook(
        folder_name="BIT",
        output_master_name="BIT_Faculty_Master.xlsx",
        output_alias_name="BIT_All_Campuses_Faculty.xlsx"
    )
    
    # 2. BITS
    build_consolidated_workbook(
        folder_name="BITS",
        output_master_name="BITS_Faculty_Master.xlsx",
        output_alias_name="BITS_All_Campuses_Faculty.xlsx"
    )
    
    # 3. NITs
    build_consolidated_workbook(
        folder_name="NITs",
        output_master_name="NITs_Faculty_Master.xlsx",
        output_alias_name="NITs_All_Campuses_Faculty.xlsx"
    )
    
    # 4. IIITs
    build_consolidated_workbook(
        folder_name="IIITs",
        output_master_name="IIITs_Faculty_Master.xlsx",
        output_alias_name="IIITs_All_Campuses_Faculty.xlsx"
    )
    
    print("\n=== All Consolidated Workbooks Generated Successfully ===")

if __name__ == '__main__':
    main()
