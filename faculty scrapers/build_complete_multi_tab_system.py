import os
import re
import shutil
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

STANDARD_COLUMNS = ['Name', 'Institution', 'Department', 'Designation', 'Qualification', 'Email', 'Profile URL']

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

MASTER_PATTERNS = [
    r'^(NIT|NITs|IIIT|IIITs|BIT|BITS)_(Faculty_Master|All_Campuses_Faculty|All_Branches_Faculty|By_Branch_Master|Faculty)\.xlsx$'
]

def is_master_or_alias(fname):
    for pat in MASTER_PATTERNS:
        if re.match(pat, fname, re.IGNORECASE):
            return True
    return False

def get_institution_files(folder_path):
    all_files = [f for f in os.listdir(folder_path) if f.endswith('.xlsx') and not f.startswith('~') and not f.startswith('.')]
    return sorted([f for f in all_files if not is_master_or_alias(f)])

def categorize_branch(dept_str):
    d = str(dept_str).strip()
    d_clean = re.sub(r"^(department|dept|school|centre)\s+of\s+", "", d, flags=re.IGNORECASE).strip()
    d_lower = d_clean.lower()
    
    # Check MCA / Computer Applications / CA first
    if d_lower in ["ca", "c.a.", "mca"] or re.search(r"\b(computer applications?|mca|c\.a\.|ca)\b", d_lower):
        return "MCA"
    # AI & Data Science
    if re.search(r"\b(artificial intelligence|ai &|ai and|aide|data science|dsai|machine learning)\b", d_lower):
        return "AI & DS"
    # Mathematics & Computing
    if re.search(r"\b(mathematics and computing|mathematical and computational|scientific computing)\b", d_lower):
        return "Math & Computing"
    # VLSI & Nanotechnology
    if re.search(r"\b(vlsi|nano|microelectronics)\b", d_lower):
        return "VLSI & Nano"
    # Instrumentation & Control
    if re.search(r"\b(instrumentation|control|ice|eie)\b", d_lower):
        return "ICE"
    # Mechatronics & Robotics
    if re.search(r"\b(mechatronics|robotics|automation)\b", d_lower):
        return "Mechatronics"
    # Information Technology
    if re.search(r"\b(information technology|information tech)\b", d_lower) or d_lower == "it":
        return "IT"
    # Computer Science & Engineering / Computing / Info Systems
    if re.search(r"\b(computer science|cse|computer engineering|computing|software|information systems)\b", d_lower):
        return "CSE"
    # Electronics & Communication / Telecommunication / Signal Processing
    if re.search(r"\b(electronics and communication|electronics & communication|ece|electronics and telecommunication|etc|telecommunication|signal processing)\b", d_lower):
        return "ECE"
    # Electrical & Electronics / Electrical Engineering
    if re.search(r"\b(electrical and electronics|electrical & electronics|eee|electrical engineering|electrical)\b", d_lower):
        return "EEE"
    # Pure Electronics
    if re.search(r"\b(electronics engineering|electronics)\b", d_lower):
        return "ECE"
        
    for c in ["[", "]", ":", "*", "?", "/", "\\"]:
        d_clean = d_clean.replace(c, "")
    return d_clean[:28].strip()

def get_clean_campus_tab_name(fname, df):
    if len(df) > 0 and 'Institution' in df.columns:
        inst = str(df['Institution'].iloc[0]).strip()
        if 'Maulana Azad' in inst: return 'MANIT Bhopal'
        if 'Malaviya' in inst: return 'MNIT Jaipur'
        if 'Motilal Nehru' in inst: return 'MNNIT Allahabad'
        if 'Visvesvaraya' in inst: return 'VNIT Nagpur'
        if 'Sardar Vallabhbhai' in inst or 'SVNIT' in inst: return 'SVNIT Surat'
        if 'IIEST' in inst: return 'IIEST Shibpur'
        if 'National Institute of Technology' in inst:
            clean = inst.replace('National Institute of Technology', 'NIT').strip()
            if len(clean) <= 31: return clean
        if 'Indian Institute of Information Technology' in inst:
            clean = inst.replace('Indian Institute of Information Technology', 'IIIT').strip()
            if len(clean) <= 31: return clean
        if 'K. K. Birla' in inst or 'Goa Campus' in inst: return 'BITS Goa'
        if 'Hyderabad Campus' in inst: return 'BITS Hyderabad'
        if 'Pilani Campus' in inst: return 'BITS Pilani'
        if 'Dubai Campus' in inst: return 'BITS Dubai'
        if 'Main Campus' in inst: return 'BIT Mesra'
        if 'PDPM IIITDM' in inst: return 'IIITDM Jabalpur'
        if 'ABV-IIITM' in inst: return 'IIITM Gwalior'
        if 'NIT Karnataka' in inst: return 'NIT Surathkal'
        if 'NIT Tiruchirappalli' in inst: return 'NIT Trichy'
        if 'IIIT Tiruchirappalli' in inst: return 'IIIT Trichy'
        if len(inst) <= 31 and not any(c in inst for c in ['[', ']', ':', '*', '?', '/', '\\']):
            return inst

    name = fname.replace('_Faculty.xlsx', '').replace('_Complete', '').replace('_', ' ')
    name = re.sub(r'[\[\]\:\*\?\/\\ ]+', ' ', name).strip()
    return name[:31].strip()

def style_sheet(ws):
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 26
    
    for col_idx, cell in enumerate(ws[1], 1):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
        
    col_widths = {col_idx: len(str(cell.value or '')) for col_idx, cell in enumerate(ws[1], 1)}
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        for col_idx, cell in enumerate(row, 1):
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN
            if is_even:
                cell.fill = ZEBRA_FILL
            
            val_str = str(cell.value or '')
            col_widths[col_idx] = max(col_widths[col_idx], len(val_str))
            
            col_name = ws.cell(row=1, column=col_idx).value
            if col_name in ['Qualification']:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    for col_idx, max_len in col_widths.items():
        col_letter = get_column_letter(col_idx)
        adjusted_width = min(max(max_len + 4, 12), 48)
        ws.column_dimensions[col_letter].width = adjusted_width
        
    ws.auto_filter.ref = ws.dimensions

def populate_sheet(ws, df):
    ws.append(STANDARD_COLUMNS)
    for _, row in df.iterrows():
        ws.append(row.tolist())
    style_sheet(ws)

def update_individual_institution_files(folder_name):
    folder_path = os.path.join(BASE_DIR, folder_name)
    if not os.path.exists(folder_path):
        return

    files = get_institution_files(folder_path)
    print(f"\n--- Updating Individual Files in {folder_name} ({len(files)} files) ---")
    
    for f in files:
        fpath = os.path.join(folder_path, f)
        # Always read from 'All Faculty' if multi-sheet already exists, else sheet 0
        wb_check = openpyxl.load_workbook(fpath, read_only=True)
        sheet_to_read = 'All Faculty' if 'All Faculty' in wb_check.sheetnames else wb_check.sheetnames[0]
        wb_check.close()
        
        df = pd.read_excel(fpath, sheet_name=sheet_to_read)
        for col in STANDARD_COLUMNS:
            if col not in df.columns:
                df[col] = "N/A"
        df = df[STANDARD_COLUMNS]
        
        df['Branch_Category'] = df['Department'].apply(categorize_branch)
        
        wb = openpyxl.Workbook()
        default_sheet = wb.active
        
        # Tab 1: All Faculty
        ws_all = wb.create_sheet(title="All Faculty")
        populate_sheet(ws_all, df[STANDARD_COLUMNS])
        
        # Branch tabs sorted by count
        branch_counts = df['Branch_Category'].value_counts()
        for branch, count in branch_counts.items():
            branch_df = df[df['Branch_Category'] == branch][STANDARD_COLUMNS]
            tab_title = branch[:31].strip()
            ws_branch = wb.create_sheet(title=tab_title)
            populate_sheet(ws_branch, branch_df)
            
        wb.remove(default_sheet)
        wb.save(fpath)
        tab_summary = ", ".join([f"{b}: {c}" for b, c in branch_counts.items()])
        print(f"  [Updated] {f} -> {len(df)} records | Tabs: [All Faculty, {tab_summary}]")

def build_campus_master(folder_name, primary_name, aliases):
    folder_path = os.path.join(BASE_DIR, folder_name)
    files = get_institution_files(folder_path)
    
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    
    all_rows = []
    used_tab_names = set()
    
    for f in files:
        fpath = os.path.join(folder_path, f)
        wb_check = openpyxl.load_workbook(fpath, read_only=True)
        sheet_to_read = 'All Faculty' if 'All Faculty' in wb_check.sheetnames else wb_check.sheetnames[0]
        wb_check.close()
        
        df = pd.read_excel(fpath, sheet_name=sheet_to_read)[STANDARD_COLUMNS]
        all_rows.append(df)
        
        tab_name = get_clean_campus_tab_name(f, df)
        if tab_name in used_tab_names:
            tab_name = f"{tab_name[:28]}_{len(used_tab_names)}"
        used_tab_names.add(tab_name)
        
        ws = wb.create_sheet(title=tab_name)
        populate_sheet(ws, df)

    all_df = pd.concat(all_rows, ignore_index=True) if all_rows else pd.DataFrame(columns=STANDARD_COLUMNS)
    all_ws = wb.create_sheet(title="All Faculty", index=0)
    populate_sheet(all_ws, all_df)
    
    wb.remove(default_sheet)
    
    p_path = os.path.join(folder_path, primary_name)
    wb.save(p_path)
    print(f"  [Saved Campus Master] {p_path} ({len(wb.sheetnames)} tabs, {len(all_df)} records)")
    
    for alias in aliases:
        a_path = os.path.join(folder_path, alias)
        shutil.copyfile(p_path, a_path)
        print(f"  [Copied Alias]        {a_path}")

def build_branch_master(folder_name, primary_name, aliases):
    folder_path = os.path.join(BASE_DIR, folder_name)
    files = get_institution_files(folder_path)
    
    all_rows = []
    for f in files:
        fpath = os.path.join(folder_path, f)
        wb_check = openpyxl.load_workbook(fpath, read_only=True)
        sheet_to_read = 'All Faculty' if 'All Faculty' in wb_check.sheetnames else wb_check.sheetnames[0]
        wb_check.close()
        
        df = pd.read_excel(fpath, sheet_name=sheet_to_read)[STANDARD_COLUMNS]
        all_rows.append(df)
        
    all_df = pd.concat(all_rows, ignore_index=True) if all_rows else pd.DataFrame(columns=STANDARD_COLUMNS)
    all_df['Branch_Category'] = all_df['Department'].apply(categorize_branch)
    
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    
    all_ws = wb.create_sheet(title="All Branches")
    populate_sheet(all_ws, all_df[STANDARD_COLUMNS])
    
    branch_counts = all_df['Branch_Category'].value_counts()
    for branch, count in branch_counts.items():
        branch_df = all_df[all_df['Branch_Category'] == branch][STANDARD_COLUMNS]
        tab_title = branch[:31].strip()
        ws = wb.create_sheet(title=tab_title)
        populate_sheet(ws, branch_df)
        
    wb.remove(default_sheet)
    
    p_path = os.path.join(folder_path, primary_name)
    wb.save(p_path)
    print(f"  [Saved Branch Master] {p_path} ({len(wb.sheetnames)} tabs, {len(all_df)} records)")
    
    for alias in aliases:
        a_path = os.path.join(folder_path, alias)
        shutil.copyfile(p_path, a_path)
        print(f"  [Copied Alias]        {a_path}")

def main():
    print("================================================================")
    print("BUILDING COMPLETE MULTI-TAB WORKBOOK ECOSYSTEM (CORRECTED)")
    print("================================================================")
    
    tiers = ['NITs', 'IIITs', 'BIT', 'BITS']
    
    for tier in tiers:
        update_individual_institution_files(tier)
        
    print("\n================================================================")
    print("GENERATING MASTER FILES FOR EACH TIER")
    print("================================================================")
    
    # NITs
    print("\n--- NITs Masters ---")
    build_campus_master(
        folder_name="NITs",
        primary_name="NITs_Faculty_Master.xlsx",
        aliases=["NIT_Faculty_Master.xlsx", "NITs_All_Campuses_Faculty.xlsx", "NIT_All_Campuses_Faculty.xlsx", "NIT_Faculty.xlsx", "NITs_Faculty.xlsx"]
    )
    build_branch_master(
        folder_name="NITs",
        primary_name="NITs_By_Branch_Master.xlsx",
        aliases=["NIT_By_Branch_Master.xlsx", "NITs_All_Branches_Faculty.xlsx", "NIT_All_Branches_Faculty.xlsx"]
    )
    
    # IIITs
    print("\n--- IIITs Masters ---")
    build_campus_master(
        folder_name="IIITs",
        primary_name="IIITs_Faculty_Master.xlsx",
        aliases=["IIIT_Faculty_Master.xlsx", "IIITs_All_Campuses_Faculty.xlsx", "IIIT_All_Campuses_Faculty.xlsx", "IIIT_Faculty.xlsx", "IIITs_Faculty.xlsx"]
    )
    build_branch_master(
        folder_name="IIITs",
        primary_name="IIITs_By_Branch_Master.xlsx",
        aliases=["IIIT_By_Branch_Master.xlsx", "IIITs_All_Branches_Faculty.xlsx", "IIIT_All_Branches_Faculty.xlsx"]
    )
    
    # BIT
    print("\n--- BIT Masters ---")
    build_campus_master(
        folder_name="BIT",
        primary_name="BIT_Faculty_Master.xlsx",
        aliases=["BIT_All_Campuses_Faculty.xlsx", "BIT_Faculty.xlsx"]
    )
    build_branch_master(
        folder_name="BIT",
        primary_name="BIT_By_Branch_Master.xlsx",
        aliases=["BIT_All_Branches_Faculty.xlsx"]
    )
    
    # BITS
    print("\n--- BITS Masters ---")
    build_campus_master(
        folder_name="BITS",
        primary_name="BITS_Faculty_Master.xlsx",
        aliases=["BITS_All_Campuses_Faculty.xlsx", "BITS_Faculty.xlsx"]
    )
    build_branch_master(
        folder_name="BITS",
        primary_name="BITS_By_Branch_Master.xlsx",
        aliases=["BITS_All_Branches_Faculty.xlsx"]
    )
    
    print("\n================================================================")
    print("ALL MULTI-TAB WORKBOOKS & MASTERS SUCCESSFULLY GENERATED")
    print("================================================================")

if __name__ == '__main__':
    main()
